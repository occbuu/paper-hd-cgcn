# -*- coding: utf-8 -*-
"""K-scan, leave-one-out (Bùi 2019), genre-split, and MDS for Paper4L.

Public package: precomputed outputs live in ../results/.
The underlying .docx corpus is NOT redistributed (publisher copyright).
To re-run locally, set CORPUS_DIR to a folder of the 34 .docx files, and place
vietnamese-wordlist.txt + vietnamese-stopwords.txt next to this script
(or allow first-run download from the public GitHub sources used by the
original pipeline).
"""
import os, re, unicodedata, urllib.request
from pathlib import Path
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.decomposition import NMF
from sklearn.cluster import KMeans
from sklearn.metrics import silhouette_score, adjusted_rand_score
from sklearn.metrics.pairwise import cosine_similarity
from sklearn.manifold import MDS

HERE = Path(__file__).resolve().parent
PKG = HERE.parent
OUT = PKG / "results"
OUT.mkdir(parents=True, exist_ok=True)

# Local re-run only. Not shipped in the public folder.
_DEFAULT_CORPUS = PKG.parent / "hopdong-CGCN" / "Dataset_Dinh_tinh"
SRC = Path(os.environ.get("CORPUS_DIR", _DEFAULT_CORPUS))

WORDLIST_FILE = HERE / "vietnamese-wordlist.txt"
STOPWORDS_FILE = HERE / "vietnamese-stopwords.txt"
MAX_SPAN = 5
_DOMAIN_TERMS = [
    "chuyển giao công nghệ", "sở hữu công nghiệp", "đổi mới sáng tạo",
    "bí quyết kỹ thuật", "công nghệ cao", "công nghệ tiên tiến",
    "công nghệ mới", "công nghệ sạch", "thị trường khoa học",
    "quản lý nhà nước", "dự án đầu tư", "bồi thường thiệt hại",
    "thẩm định giá", "quyền sở hữu", "việt nam", "khoa học công nghệ",
    "chuyển giao công nghệ trong nước", "vốn đầu tư công",
    "trung quốc", "hàn quốc", "nhật bản", "hoa kỳ", "châu âu", "đông nam á",
    "liên minh châu âu", "phát thải",
]
_ABBREV_PATTERNS = [
    (re.compile(r"\bkh\s*&\s*cn\b|\bkh-cn\b|\bkhcn\b"), "khoa học công nghệ"),
    (re.compile(r"\bcgcn\b"), "chuyển giao công nghệ"),
    (re.compile(r"\bdnnn\b"), "doanh nghiệp nhà nước"),
    (re.compile(r"\bdn\b"), "doanh nghiệp"),
    (re.compile(r"\bshtt\b"), "sở hữu trí tuệ"),
    (re.compile(r"\bubnd\b"), "ủy ban nhân dân"),
]
_SOURCES = {
    WORDLIST_FILE: "https://raw.githubusercontent.com/duyet/vietnamese-wordlist/master/Viet74K.txt",
    STOPWORDS_FILE: "https://raw.githubusercontent.com/stopwords/vietnamese-stopwords/master/vietnamese-stopwords.txt",
}


def _clean_wordlist(raw_text: str) -> str:
    words = set()
    for line in raw_text.splitlines():
        w = unicodedata.normalize("NFC", line.strip().lower())
        if w and all(c.isalpha() or c == " " for c in w):
            words.add(w)
    for t in _DOMAIN_TERMS:
        words.add(unicodedata.normalize("NFC", t.lower()))
    return "\n".join(sorted(words)) + "\n"


def ensure_lexicon():
    for path, url in _SOURCES.items():
        if path.exists():
            continue
        print(f"  [download] {path.name}")
        with urllib.request.urlopen(url, timeout=30) as resp:
            raw = resp.read().decode("utf-8")
        if path == WORDLIST_FILE:
            raw = _clean_wordlist(raw)
        path.write_text(raw, encoding="utf-8")


def load_wordset(path):
    return {w.strip() for w in open(path, encoding="utf-8") if w.strip()}


def segment_words(sylls, wordset, max_span=MAX_SPAN):
    out, i, n = [], 0, len(sylls)
    while i < n:
        matched = False
        for span in range(min(max_span, n - i), 1, -1):
            cand = " ".join(sylls[i:i + span])
            if cand in wordset:
                out.append(cand.replace(" ", "_"))
                i += span
                matched = True
                break
        if not matched:
            out.append(sylls[i])
            i += 1
    return out


def _normalize_text(s: str) -> str:
    s = unicodedata.normalize("NFC", s.lower())
    s = re.sub(r"https?://\S+|www\.\S+", " ", s)
    for pat, repl in _ABBREV_PATTERNS:
        s = pat.sub(repl, s)
    return s


def preprocess(s: str, vn_words, stopwords) -> str:
    s = _normalize_text(s)
    s = re.sub(r"[^\w\s]", " ", s, flags=re.UNICODE)
    s = re.sub(r"\d+", " ", s)
    sylls = s.split()
    tokens = segment_words(sylls, vn_words) if vn_words else sylls
    return " ".join(t for t in tokens if t.replace("_", " ") not in stopwords and len(t) > 1)


def read_docx(path: str) -> str:
    from docx import Document
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)


def load_corpus(vn_words, stopwords):
    if not SRC.is_dir():
        raise SystemExit(
            f"Corpus folder not found: {SRC}\n"
            "This public package ships precomputed files in results/.\n"
            "To re-run, set CORPUS_DIR to a directory of the 34 .docx files."
        )
    paths = sorted(p for p in SRC.glob("*.docx") if not p.name.startswith("~$"))
    rows = []
    for i, path in enumerate(paths):
        rows.append({
            "id": f"Bài {i+1}",
            "filename": path.name,
            "text": read_docx(str(path)).strip(),
        })
    df = pd.DataFrame(rows)
    legal_prefix = df["filename"].str.startswith("HDCG_") | df["filename"].str.contains(
        "minh-dangthu", case=False
    )
    df["genre"] = np.where(legal_prefix, "legal_scholarly", "news_policy")
    df["is_bui2019"] = df["filename"].str.contains("HDCG_A09", regex=False)
    df["clean"] = df["text"].apply(lambda s: preprocess(s, vn_words, stopwords))
    return df


def fit_nmf(texts, k, random_state=42):
    tfidf = TfidfVectorizer(max_df=0.9, min_df=2)
    X = tfidf.fit_transform(texts)
    model = NMF(n_components=k, random_state=random_state, max_iter=500, init="nndsvda")
    W = model.fit_transform(X)
    labels = W.argmax(axis=1)
    terms = tfidf.get_feature_names_out()
    keywords = []
    for j in range(k):
        top = model.components_[j].argsort()[::-1][:8]
        keywords.append(", ".join(terms[i] for i in top))
    sil = silhouette_score(X, labels, metric="cosine") if len(set(labels)) >= 2 else np.nan
    max_share = pd.Series(labels).value_counts().max() / len(labels)
    km = KMeans(n_clusters=k, random_state=random_state, n_init=10).fit_predict(X)
    ari = adjusted_rand_score(labels, km)
    return dict(X=X, W=W, H=model.components_, labels=labels, keywords=keywords,
                silhouette=sil, max_share=max_share, ari_kmeans=ari, tfidf=tfidf)


def topic_table(df, labels, keywords, k):
    rows = []
    for j in range(k):
        idx = np.where(labels == j)[0]
        rows.append({
            "topic": f"NMF-{j+1}",
            "n": int(len(idx)),
            "keywords": keywords[j],
            "docs": "; ".join(df.iloc[idx]["id"].tolist()),
            "files": "; ".join(df.iloc[idx]["filename"].tolist()),
        })
    return pd.DataFrame(rows)


def main():
    ensure_lexicon()
    vn_words = load_wordset(WORDLIST_FILE)
    vn_words |= {unicodedata.normalize("NFC", t.lower()) for t in _DOMAIN_TERMS}
    stopwords = load_wordset(STOPWORDS_FILE)
    stopwords |= {"quy định", "điều", "khoản", "luật", "chương", "mục",
                  "nghị định", "thông tư", "số", "ngày", "tháng", "năm",
                  "ts", "ths", "pgs", "gs", "th.s",
                  "image", "javascript", "css", "script", "trang"}

    df = load_corpus(vn_words, stopwords)
    df[["id", "filename", "genre", "is_bui2019"]].to_csv(
        OUT / "corpus_inventory.csv", index=False, encoding="utf-8-sig")
    print("n =", len(df),
          "legal_scholarly =", int((df.genre == "legal_scholarly").sum()),
          "news_policy =", int((df.genre == "news_policy").sum()))

    scan = []
    fits = {}
    for k in range(3, 9):
        fit = fit_nmf(df["clean"], k)
        fits[k] = fit
        scan.append({
            "K": k,
            "silhouette": round(float(fit["silhouette"]), 3),
            "largest_topic_share_%": round(100 * float(fit["max_share"]), 1),
            "ARI_NMF_KMeansTFIDF": round(float(fit["ari_kmeans"]), 3),
            "n_distinct_topics": int(len(set(fit["labels"]))),
        })
        print("K-scan", scan[-1])
    scan_df = pd.DataFrame(scan)
    scan_df.to_csv(OUT / "k_scan.csv", index=False, encoding="utf-8-sig")

    fig, ax = plt.subplots(figsize=(7.6, 4.2))
    ax.plot(scan_df["K"], scan_df["silhouette"], marker="o", label="Silhouette")
    ax.plot(scan_df["K"], scan_df["ARI_NMF_KMeansTFIDF"], marker="s",
            label="ARI (NMF, K-means)")
    ax.set_xlabel("K")
    ax.set_ylabel("Score")
    ax.set_xticks(list(range(3, 9)))
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.legend(frameon=False)
    ax.set_title("Topic-number scan on the 34-document corpus")
    fig.tight_layout()
    fig.savefig(OUT / "fig_k_scan.png", dpi=180)
    plt.close()

    fit7 = fits[7]
    df["nmf7"] = [f"NMF-{i+1}" for i in fit7["labels"]]
    topic_table(df, fit7["labels"], fit7["keywords"], 7).to_csv(
        OUT / "topics_k7.csv", index=False, encoding="utf-8-sig")
    df[["id", "filename", "genre", "nmf7"]].to_csv(
        OUT / "doc_assignment_k7.csv", index=False, encoding="utf-8-sig")

    dist = np.clip(1 - cosine_similarity(fit7["X"]), 0, None)
    coords = MDS(n_components=2, dissimilarity="precomputed", random_state=42,
                 n_init=4, normalized_stress="auto").fit_transform(dist)
    cmap = plt.cm.tab10
    plt.figure(figsize=(9, 7))
    for j in range(7):
        name = f"NMF-{j+1}"
        idx = df["nmf7"] == name
        plt.scatter(coords[idx, 0], coords[idx, 1], s=48, color=cmap(j / 6), label=name)
    plt.legend(fontsize=8, loc="center left", bbox_to_anchor=(1.0, 0.5))
    plt.title("MDS of the 34 documents (TF-IDF geometry; colour = NMF, K = 7)")
    plt.tight_layout()
    plt.savefig(OUT / "fig_mds.png", dpi=180)
    plt.close()

    mask = ~df["is_bui2019"]
    fit_loo = fit_nmf(df.loc[mask, "clean"], 7)
    loo_df = df.loc[mask].copy()
    loo_df["nmf7_loo"] = [f"NMF-{i+1}" for i in fit_loo["labels"]]
    loo_topics = topic_table(loo_df, fit_loo["labels"], fit_loo["keywords"], 7)
    loo_topics.to_csv(OUT / "topics_leave_one_out_bui2019.csv", index=False, encoding="utf-8-sig")
    loo_df[["id", "filename", "genre", "nmf7", "nmf7_loo"]].to_csv(
        OUT / "doc_assignment_leave_one_out_bui2019.csv", index=False, encoding="utf-8-sig")

    news = df[df["genre"] == "news_policy"].copy()
    fit_news = fit_nmf(news["clean"], 7)
    news_topics = topic_table(news, fit_news["labels"], fit_news["keywords"], 7)
    news_topics.to_csv(OUT / "topics_news_policy_only.csv", index=False, encoding="utf-8-sig")

    legal = df[df["genre"] == "legal_scholarly"].copy()
    fit_legal = fit_nmf(legal["clean"], 3)
    legal_topics = topic_table(legal, fit_legal["labels"], fit_legal["keywords"], 3)
    legal_topics.to_csv(OUT / "topics_legal_scholarly_only.csv", index=False, encoding="utf-8-sig")

    with open(OUT / "robustness_summary.txt", "w", encoding="utf-8") as f:
        f.write(scan_df.to_string(index=False))
        f.write("\n\n=== K=7 full corpus ===\n")
        f.write(topic_table(df, fit7["labels"], fit7["keywords"], 7)[["topic", "n", "keywords"]].to_string(index=False))
        f.write("\n\n=== Leave-one-out Bùi 2019 ===\n")
        f.write(loo_topics[["topic", "n", "keywords"]].to_string(index=False))
        f.write("\n\n=== News/policy only K=7 ===\n")
        f.write(news_topics[["topic", "n", "keywords"]].to_string(index=False))
        f.write("\n\n=== Legal/scholarly only K=3 ===\n")
        f.write(legal_topics[["topic", "n", "keywords"]].to_string(index=False))
    print("wrote", OUT)


if __name__ == "__main__":
    main()
